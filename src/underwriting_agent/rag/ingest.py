"""Ingest underwriting policy sources into the shared RagPipeline vector index.

Chunks are tagged with a ``program`` metadata field so retrieval / multi-doc
reasoning can filter by layer instead of treating all policy as one pool:

- ``compliance_floor`` — ECOA / Reg B (boolean gate; never scored)
- ``eligibility_gate`` — SBA core eligibility adopted as universal min bar
- ``sba_7a`` — SBA 7(a) program underwriting
- ``cdfi_direct`` — Accion-style direct CDFI product criteria
"""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from langchain_core.documents import Document

from underwriting_agent.models import PolicyLayer
from underwriting_agent.rag import POLICY_INDEX, get_policy_pipeline, policy_sources_dir
from underwriting_agent.rag.tags import resolve_program

logger = logging.getLogger(__name__)


def _base_metadata(path: Path, *, page: int, program: PolicyLayer) -> dict[str, Any]:
    return {
        "page": page,
        "source": path.name,
        "file": path.name,
        "created_at": datetime.now().isoformat(),
        "corpus": "underwriting_policy",
        "program": program.value,
    }


def _load_pdf_pages(path: Path) -> list[Document]:
    import pdf_inspector

    result = pdf_inspector.extract_pages_markdown(path.as_posix())
    docs: list[Document] = []
    for page in result.pages:
        text = (page.markdown or "").strip()
        if not text:
            continue
        program = resolve_program(path.name, text)
        docs.append(
            Document(
                page_content=text,
                metadata=_base_metadata(path, page=page.page, program=program),
            )
        )
    return docs


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(path.as_posix())
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    if not paragraphs:
        return []
    docs: list[Document] = []
    bucket: list[str] = []
    chunk_idx = 0
    for para in paragraphs:
        bucket.append(para)
        if len(bucket) >= 8:
            text = "\n\n".join(bucket)
            program = resolve_program(path.name, text)
            docs.append(
                Document(
                    page_content=text,
                    metadata=_base_metadata(path, page=chunk_idx, program=program),
                )
            )
            bucket = []
            chunk_idx += 1
    if bucket:
        text = "\n\n".join(bucket)
        program = resolve_program(path.name, text)
        docs.append(
            Document(
                page_content=text,
                metadata=_base_metadata(path, page=chunk_idx, program=program),
            )
        )
    return docs


def _load_markdown(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    # Split on ## headings so criteria sections stay retrieval-friendly.
    parts = re.split(r"(?=^##\s)", text, flags=re.M)
    docs: list[Document] = []
    for i, part in enumerate(parts):
        chunk = part.strip()
        if not chunk:
            continue
        program = resolve_program(path.name, chunk)
        docs.append(
            Document(
                page_content=chunk,
                metadata=_base_metadata(path, page=i, program=program),
            )
        )
    return docs


def load_policy_documents(data_dir: Path | None = None) -> list[Document]:
    """Load policy files from ``data/policy_sources`` (skips SOURCES.md / .gitkeep)."""
    root = data_dir or policy_sources_dir()
    documents: list[Document] = []
    for path in sorted(root.iterdir()):
        if path.name.startswith(".") or path.name.upper().startswith("SOURCES"):
            continue
        if path.suffix.lower() in {".gitkeep"} or not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            documents.extend(_load_pdf_pages(path))
        elif suffix in {".docx", ".doc"}:
            documents.extend(_load_docx(path))
        elif suffix == ".md":
            documents.extend(_load_markdown(path))
        else:
            logger.warning("Skipping unsupported policy file: %s", path.name)
    by_program: dict[str, int] = {}
    for d in documents:
        key = str(d.metadata.get("program", "unknown"))
        by_program[key] = by_program.get(key, 0) + 1
    logger.info(
        "Loaded %s source units from %s (%s)",
        len(documents),
        root,
        ", ".join(f"{k}={v}" for k, v in sorted(by_program.items())),
    )
    return documents


def ingest_policy_sources(*, targets: tuple[str, ...] = ("vector",)) -> int:
    """Chunk + embed policy docs into the underwriting pgvector index."""
    docs = load_policy_documents()
    if not docs:
        raise FileNotFoundError(f"No policy documents found in {policy_sources_dir()}")
    pipeline = get_policy_pipeline(index_name=POLICY_INDEX)
    pipeline.ingest(docs, index_name=POLICY_INDEX, targets=targets)
    return len(docs)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    n = ingest_policy_sources()
    print(f"Ingested {n} source units into index {POLICY_INDEX!r}")


if __name__ == "__main__":
    main()
