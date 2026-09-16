"""Ingest underwriting policy sources into the shared RagPipeline vector index.

Chunks are tagged with ``program`` (policy layer / product type) and optional
``lender_id`` (lender overlay) so retrieval can filter shared rules vs
lender-specific criteria:

- ``compliance_floor`` — ECOA / Reg B (boolean gate; never scored)
- ``eligibility_gate`` — SBA core eligibility adopted as universal min bar
- ``sba_7a`` — SBA 7(a) program underwriting (generic; ``lender_id`` null)
- ``cdfi_direct`` — CDFI Direct product underwriting (generic when no lender)
- Lender overlays (e.g. Accion SBA 7(a)) set ``lender_id`` and keep ``program``
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from langchain_core.documents import Document

from config import IngestTarget, get_settings
from models import PolicyLayer
from policy_rag import get_policy_pipeline, policy_sources_dir
from agents.lender_config import upsert_lender_config_from_text
from policy_rag.catalog import lookup_source
from policy_rag.tags import resolve_program

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class IngestResult:
    """Outcome of a policy corpus rebuild."""

    index_name: str
    targets: tuple[str, ...]
    source_units: int
    by_program: dict[str, int] = field(default_factory=dict)
    source_files: list[str] = field(default_factory=list)


def _base_metadata(path: Path, *, page: int, program: PolicyLayer) -> dict[str, Any]:
    entry = lookup_source(path.name)
    meta: dict[str, Any] = {
        "page": page,
        "source": path.name,
        "file": path.name,
        "created_at": datetime.now().isoformat(),
        "corpus": "underwriting_policy",
        "program": program.value,
    }
    if entry is None:
        return meta
    meta["title"] = entry.title
    meta["authority"] = entry.authority
    if entry.lender_id:
        meta["lender_id"] = entry.lender_id
    if entry.url:
        meta["url"] = entry.url
    if entry.version:
        meta["version"] = entry.version
    if entry.effective_date:
        meta["effective_date"] = entry.effective_date
    return meta


def _load_pdf_pages(path: Path) -> list[Document]:
    import pdf_inspector

    result = pdf_inspector.extract_pages_markdown(path.as_posix())
    docs: list[Document] = []
    for page in result.pages:
        text = (page.markdown or "").strip()
        if not text:
            continue
        program = resolve_program(path.name, text)
        docs.append(
            Document(
                page_content=text,
                metadata=_base_metadata(path, page=page.page, program=program),
            )
        )
    return docs


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(path.as_posix())
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    if not paragraphs:
        return []
    docs: list[Document] = []
    bucket: list[str] = []
    chunk_idx = 0
    for para in paragraphs:
        bucket.append(para)
        if len(bucket) >= 8:
            text = "\n\n".join(bucket)
            program = resolve_program(path.name, text)
            docs.append(
                Document(
                    page_content=text,
                    metadata=_base_metadata(path, page=chunk_idx, program=program),
                )
            )
            bucket = []
            chunk_idx += 1
    if bucket:
        text = "\n\n".join(bucket)
        program = resolve_program(path.name, text)
        docs.append(
            Document(
                page_content=text,
                metadata=_base_metadata(path, page=chunk_idx, program=program),
            )
        )
    return docs


def _load_markdown(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    # Split on ## headings so criteria sections stay retrieval-friendly.
    parts = re.split(r"(?=^##\s)", text, flags=re.M)
    docs: list[Document] = []
    for i, part in enumerate(parts):
        chunk = part.strip()
        if not chunk:
            continue
        program = resolve_program(path.name, chunk)
        docs.append(
            Document(
                page_content=chunk,
                metadata=_base_metadata(path, page=i, program=program),
            )
        )
    return docs


def list_policy_source_files(data_dir: Path | None = None) -> list[str]:
    """Return ingestible filenames under the policy sources directory."""
    root = data_dir or policy_sources_dir()
    if not root.is_dir():
        return []
    names: list[str] = []
    for path in sorted(root.iterdir()):
        if path.name.startswith(".") or path.name.upper().startswith("SOURCES"):
            continue
        if path.name == "catalog.json":
            continue
        if path.suffix.lower() in {".gitkeep"} or not path.is_file():
            continue
        if path.suffix.lower() in {".pdf", ".docx", ".doc", ".md"}:
            names.append(path.name)
    return names


def load_policy_documents(data_dir: Path | None = None) -> list[Document]:
    """Load policy files from ``data/policy_sources`` (skips SOURCES.md / .gitkeep).

    For lender-tagged sources, also extracts structured front matter into
    ``data/lenders/{lender_id}.json`` + refreshes ``registry.json``.
    """
    root = data_dir or policy_sources_dir()
    documents: list[Document] = []
    lender_texts: dict[str, list[str]] = {}
    for path in sorted(root.iterdir()):
        if path.name.startswith(".") or path.name.upper().startswith("SOURCES"):
            continue
        if path.name == "catalog.json":
            continue
        if path.suffix.lower() in {".gitkeep"} or not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            docs = _load_pdf_pages(path)
        elif suffix in {".docx", ".doc"}:
            docs = _load_docx(path)
        elif suffix == ".md":
            docs = _load_markdown(path)
        else:
            logger.warning("Skipping unsupported policy file: %s", path.name)
            continue
        documents.extend(docs)
        entry = lookup_source(path.name)
        if entry is not None and entry.lender_id:
            lender_texts.setdefault(path.name, []).extend(
                d.page_content for d in docs if d.page_content.strip()
            )

    for source_name, parts in lender_texts.items():
        entry = lookup_source(source_name)
        if entry is None or not entry.lender_id:
            continue
        upsert_lender_config_from_text(
            text="\n\n".join(parts),
            source_file=source_name,
            catalog_lender_id=entry.lender_id,
            catalog_program=entry.program.value,
            catalog_url=entry.url,
        )

    by_program: dict[str, int] = {}
    for d in documents:
        key = str(d.metadata.get("program", "unknown"))
        by_program[key] = by_program.get(key, 0) + 1
    logger.info(
        "Loaded %s source units from %s (%s)",
        len(documents),
        root,
        ", ".join(f"{k}={v}" for k, v in sorted(by_program.items())),
    )
    return documents


def _program_counts(docs: list[Document]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for d in docs:
        key = str(d.metadata.get("program", "unknown"))
        counts[key] = counts.get(key, 0) + 1
    return counts


def ingest_policy_sources(
    *,
    targets: tuple[IngestTarget, ...] | None = None,
    index_name: str | None = None,
    data_dir: Path | None = None,
) -> IngestResult:
    """Chunk + embed policy docs into the configured index (default: pgvector only)."""
    settings = get_settings()
    resolved_targets = targets if targets is not None else settings.ingest_targets()
    resolved_index = index_name or settings.rag_index_name
    source_files = list_policy_source_files(data_dir)
    docs = load_policy_documents(data_dir)
    if not docs:
        raise FileNotFoundError(f"No policy documents found in {data_dir or policy_sources_dir()}")
    pipeline = get_policy_pipeline(index_name=resolved_index)
    pipeline.ingest(docs, index_name=resolved_index, targets=resolved_targets)
    try:
        from agents.lenders import reload_lenders

        reload_lenders()
    except Exception:  # noqa: BLE001
        logger.debug("reload_lenders skipped", exc_info=True)
    return IngestResult(
        index_name=resolved_index,
        targets=tuple(resolved_targets),
        source_units=len(docs),
        by_program=_program_counts(docs),
        source_files=source_files,
    )


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    result = ingest_policy_sources()
    print(
        f"Ingested {result.source_units} source units into index "
        f"{result.index_name!r} targets={list(result.targets)}"
    )


if __name__ == "__main__":
    main()
